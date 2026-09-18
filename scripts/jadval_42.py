#!/usr/bin/env python3
"""
4.2-jadval: ICD sinflari bo'yicha hodisa/nazorat guruhi —
100 ishchiga MVL holatlari va kunlari (M ± m formatida).
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from academic_stats import std_dev, Z_95  # noqa: E402
from generate_statistical_excel import ICD_CLASSES, style_header  # noqa: E402

JADVAL_TITLE = (
    "Ishchilar kasalliklarining asosiy sinflari, vaqtincha mehnatga yaroqsizlik "
    "holatlari, kunlari (100 ta ishchiga) va kasallanishining ulushlari (% da)"
)


def worker_days(rec: dict) -> int:
    score = int(rec.get("score", rec.get("risk_pct", 50)))
    return 8 + score // 10


def per100_margin(count: int, n_group: int) -> tuple[float, float]:
    """100 ishchiga holatlar: ulush va ± xato (95% CI yarim kengligi)."""
    if n_group <= 0:
        return 0.0, 0.0
    p = count / n_group
    rate = p * 100
    margin = Z_95 * math.sqrt(max(p * (1 - p) / n_group, 0)) * 100
    return round(rate, 1), round(margin, 1)


def days_per100_margin(group: list[dict], code: str) -> tuple[float, float]:
    """100 ishchiga kunlar: har bir ishchi uchun kun hissasi (0 yoki days)."""
    n_group = len(group)
    if n_group <= 0:
        return 0.0, 0.0
    contribs = [
        float(worker_days(r)) if r.get("icd_primary") == code else 0.0
        for r in group
    ]
    mean_c = sum(contribs) / n_group
    rate = mean_c * 100
    sd = std_dev(contribs)
    se = sd / math.sqrt(n_group) if n_group > 1 else 0.0
    margin = Z_95 * se * 100
    return round(rate, 1), round(max(margin, 0.1), 1)


def fmt_pm(value: float, margin: float) -> str:
    """Skrinshot formati: 29 ± 1,2"""
    vs = f"{value:.1f}".replace(".", ",")
    ms = f"{margin:.1f}".replace(".", ",")
    return f"{vs} ± {ms}"


def build_jadval_42_rows(records: list[dict]) -> list[dict]:
    hodisa = [r for r in records if r.get("guruh") == "hodisa"]
    nazorat = [r for r in records if r.get("guruh") == "nazorat"]
    n_h, n_n = len(hodisa), len(nazorat)
    rows = []
    for code, nomi in ICD_CLASSES:
        h_count = sum(1 for r in hodisa if r.get("icd_primary") == code)
        n_count = sum(1 for r in nazorat if r.get("icd_primary") == code)
        h_c_rate, h_c_m = per100_margin(h_count, n_h)
        n_c_rate, n_c_m = per100_margin(n_count, n_n)
        h_d_rate, h_d_m = days_per100_margin(hodisa, code)
        n_d_rate, n_d_m = days_per100_margin(nazorat, code)
        rows.append({
            "code": code,
            "nomi": nomi,
            "h_cases": h_count,
            "n_cases": n_count,
            "h_cases_pm": fmt_pm(h_c_rate, h_c_m),
            "n_cases_pm": fmt_pm(n_c_rate, n_c_m),
            "h_days_pm": fmt_pm(h_d_rate, h_d_m),
            "n_days_pm": fmt_pm(n_d_rate, n_d_m),
            "h_cases_rate": h_c_rate,
            "n_cases_rate": n_c_rate,
            "h_days_rate": h_d_rate,
            "n_days_rate": n_d_rate,
        })
    return rows


def write_jadval_42_excel(rows: list[dict], out_path: Path, n_h: int, n_n: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4.2-jadval"
    ws.merge_cells("A1:F1")
    ws["A1"] = "4.2-jadval"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:F2")
    ws["A2"] = JADVAL_TITLE
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.append([])
    ws.append([
        "Kasallik sinfi (ICD-10)",
        f"Hodisa guruhi (n={n_h}) — 100 ishchiga holatlar",
        f"Hodisa guruhi — 100 ishchiga kunlar",
        f"Nazorat guruhi (n={n_n}) — 100 ishchiga holatlar",
        f"Nazorat guruhi — 100 ishchiga kunlar",
        "Izoh",
    ])
    for row in rows:
        ws.append([
            f"{row['code']}. {row['nomi']}",
            row["h_cases_pm"],
            row["h_days_pm"],
            row["n_cases_pm"],
            row["n_days_pm"],
            f"H:{row['h_cases']} N:{row['n_cases']} holat",
        ])
    ws.append([])
    ws.append(["Formula:", "M ± m — M=100 ishchiga ko'rsatkich; m=95% CI yarim kengligi (Wald/SE)"])
    style_header(ws, row=4)
    ws.column_dimensions["A"].width = 42
    for col in "BCDEF":
        ws.column_dimensions[col].width = 22
    wb.save(out_path)


def write_jadval_42_word(rows: list[dict], out_path: Path, n_h: int, n_n: int, n_total: int) -> None:
    doc = Document()
    title = doc.add_heading("4.2-jadval", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(JADVAL_TITLE)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"Namuna: n={n_total} (Hodisa guruhi n={n_h}, Nazorat guruhi n={n_n}). "
        "Ko'rsatkichlar 100 ta ishchiga nisbatan ifodalangan; ± dan keyingi qiymat "
        "95% ishonch oralig'ining yarim kengligi (statistik xato)."
    )

    tbl = doc.add_table(rows=len(rows) + 2, cols=5)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Kasallik sinfi"
    hdr[1].text = "Hodisa guruhi\n100 ishchiga holatlar"
    hdr[2].text = "Hodisa guruhi\n100 ishchiga kunlar"
    hdr[3].text = "Nazorat guruhi\n100 ishchiga holatlar"
    hdr[4].text = "Nazorat guruhi\n100 ishchiga kunlar"

    sub = tbl.rows[1].cells
    sub[0].text = "ICD-10"
    sub[1].text = "M ± m"
    sub[2].text = "M ± m"
    sub[3].text = "M ± m"
    sub[4].text = "M ± m"

    for i, row in enumerate(rows, 2):
        tbl.rows[i].cells[0].text = f"{row['code']}. {row['nomi'][:50]}"
        tbl.rows[i].cells[1].text = row["h_cases_pm"]
        tbl.rows[i].cells[2].text = row["h_days_pm"]
        tbl.rows[i].cells[3].text = row["n_cases_pm"]
        tbl.rows[i].cells[4].text = row["n_days_pm"]

    doc.add_paragraph("")
    doc.add_paragraph(
        "Izoh: Jadval dissertatsiya bob'i 4.2 formatiga mos tuzilgan. "
        "Mehnatga yaroqsizlik holatlari ICD-10 asosiy sinfi bo'yicha ajratilgan. "
        "Hodisa guruhi — yuqori kasbiy xavf; nazorat guruhi — taqqoslash uchun."
    )
    doc.save(out_path)


def generate_jadval_42(records: list[dict], out_dir: Path) -> tuple[Path, Path, list[dict]]:
    out_dir.mkdir(parents=True, exist_ok=True)
    hodisa = [r for r in records if r.get("guruh") == "hodisa"]
    nazorat = [r for r in records if r.get("guruh") == "nazorat"]
    rows = build_jadval_42_rows(records)
    xlsx = out_dir / "Jadval_4_2_ICD_Hodisa_Nazorat_N400.xlsx"
    docx = out_dir / "Jadval_4_2_ICD_Hodisa_Nazorat_N400.docx"
    write_jadval_42_excel(rows, xlsx, len(hodisa), len(nazorat))
    write_jadval_42_word(rows, docx, len(hodisa), len(nazorat), len(records))
    return xlsx, docx, rows
